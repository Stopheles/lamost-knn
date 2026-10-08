# -*- coding: utf-8 -*-
"""生成不超过2页的实验摘要 docx。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))

doc = Document()

# 全局字体
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(9.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)


def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(14); r.font.bold = True
    r.font.name = "Times New Roman"
    r.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def h(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(10.5); r.font.bold = True
    r.font.name = "Times New Roman"
    r.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)


def body(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(9.5); r.font.bold = bold
    r.font.name = "Times New Roman"
    p.paragraph_format.space_after = Pt(2)
    return p


def add_table(headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, htext in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(htext)
        r.font.bold = True; r.font.size = Pt(8.5)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(v))
            r.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


title("LAMOST 光谱 Star/Galaxy/QSO 三分类：KNN 分类实验")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("第7次课课后实践作业 · 实验摘要（含分工与复现说明）")
r.font.size = Pt(9)

h("1　任务")
body("用监督学习完成 LAMOST 光谱的恒星（Star）、星系（Galaxy）、类星体（QSO）三分类；"
     "从零实现 K 近邻（KNN）分类器，掌握标准化、距离度量、K 值选择与近邻投票；"
     "从指标、混淆矩阵和错误样本解释模型适用边界。")

h("2　数据与实验设置")
body("数据沿用第6讲作业3.1的相同样本范围与16个特征（4个颜色、4个组合颜色、5个log信噪比、"
     "红移z、z_err、Gaia G星等）。类别统一为 Star/Galaxy/QSO（剔除UNKNOWN 66,043条）；"
     "原始样本2,474,405条（Star占99.75%），STAR随机下采样至20,000，Galaxy 4,930、QSO 1,232全保留，"
     "共26,162条。分层抽样划分 train/val/test = 70%/15%/15%，随机种子42，三类比例在各子集保持一致；"
     "标准化参数仅用训练集估计，验证/测试集只做变换，无数据泄漏。测试集只在最终评估使用一次。"
     "模型均用 numpy 从零实现，并与 sklearn 预测逐点比对（一致率≥99%）验证正确性。")

h("3　方法")
body("KNN：暴力搜索近邻（批量矩阵实现），候选K∈{1,3,5,7,9,15,21}，验证集上以Macro F1为主、"
     "Accuracy为辅选择K；对照欧氏/曼哈顿距离、均匀/距离加权投票、标准化/不标准化。")

h("4　主要结果（测试集，仅评估一次）")
add_table(
    ["模型（K=5）", "Accuracy", "Macro F1", "Galaxy P/R/F1", "QSO P/R/F1"],
    [
        ["KNN 主配置(标准化+欧氏+均匀)", "0.894", "0.796", "0.77/0.74/0.76", "0.83/0.60/0.69"],
        ["KNN 距离加权", "0.896", "0.798", "0.78/0.74/0.76", "0.83/0.60/0.70"],
        ["KNN 曼哈顿距离", "0.907", "0.815", "0.80/0.78/0.79", "0.84/0.61/0.71"],
        ["KNN 未标准化", "0.960", "0.897", "0.92/0.92/0.92", "0.92/0.70/0.79"],
    ])
_pic = doc.add_paragraph(); _pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
_pic.add_run().add_picture(os.path.join(BASE, 'figures', 'test_comparison.png'), width=Cm(13.5))
_cap = doc.add_paragraph(); _cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
_cr = _cap.add_run('图1　各 KNN 配置在统一测试集上的表现')
_cr.font.size = Pt(8.5)
body("验证集选K：Macro F1在K=5达峰值0.777（K=1时0.742，K=21时0.761），稳定区间K∈[5,7]。"
     "K过小对噪声敏感（高方差），K过大邻居混入异类、边界过度平滑。")

h("5　结论与讨论")
body("（1）四种 KNN 配置在统一测试集上 Accuracy 0.894-0.960、Macro F1 0.796-0.897；"
     "其中未标准化配置最高（0.960/0.897），主配置（标准化+欧氏+均匀）为 0.894/0.796。"
     "KNN 不做分布假设，对特征相关与多峰结构较稳健。")
body("（2）标准化在本数据集上反而降低KNN表现（0.894→不标准化0.960）：区分度最强的特征"
     "（z、星等、log信噪比，第6讲互信息排序靠前）恰好方差最大，原始距离自然赋予其更大权重；"
     "等方差标准化稀释了主要信号。但“只用训练集统计量”的防泄漏流程必须严格遵守。")
body("（3）主要错分方向为Galaxy→Star（180条）与Star→Galaxy（136条，测光颜色简并）以及QSO→Star（51条，低红移QSO难辨）；"
     "QSO召回率仅0.60。低信噪比半区错误率16.9%，为高信噪比半区（4.2%）的4倍；"
     "抽查发现部分标为QSO的样本z<0.12且z_err缺失（-9999），属可疑标签噪声。")
body("（4）近邻类别比例不能直接视为可信概率：K=5时它只能取0.4/0.6/0.8/1.0四个离散值，"
     "且得票0.4的样本实际正确率约0.62，系统性偏离真实概率；它未考虑邻居距离、局部密度与类别先验，"
     "仅可作粗粒度排序参考。")

h("6　分工与复现")
body("分工：数据准备与特征核对（成员A）；KNN实现与K选择实验（成员B）；对照实验与指标汇总（成员C）；"
     "错误样本分析、图表与汇报（成员D）。（请按组内实际情况填写姓名）")
body("复现：python prep_data.py 生成 data/splits/ 划分 → 运行 knn_homework.ipynb（或 run_experiments.py + "
     "error_analysis.py）→ 输出 results/ 结果表与 figures/ 图。环境：Python 3 + numpy/pandas/matplotlib"
     "（sklearn仅用于正确性校验）；随机种子42；划分索引与参数存于 data/splits/manifest.json 与 results/summary.json。")

out = os.path.join(BASE, "实验摘要.docx")
doc.save(out)
print("saved:", out)
